"""Secure, source-agnostic source adapters. Credentials stay outside Git."""
from __future__ import annotations
import csv,os,re
from pathlib import Path
from urllib.parse import urlparse
import requests
SUPPORTED_EXTENSIONS={".csv",".xlsx",".xlsm",".json",".jsonl",".parquet"}

def safe_name(name:str)->str:
    return re.sub(r"[^A-Za-z0-9_.-]+","_",name.strip()).strip("._") or "source"

def _headers(spec:dict)->dict:
    env=spec.get("auth_env")
    if not env:return {}
    token=os.getenv(env)
    if not token:raise RuntimeError(f"Authentication secret {env} is not configured.")
    return {spec.get("auth_header","Authorization"):token}

def save_upload(uploaded,destination:Path)->Path:
    destination.parent.mkdir(parents=True,exist_ok=True)
    destination.write_bytes(uploaded.getvalue())
    return destination

def excel_to_csv(source:Path,destination_dir:Path,name:str)->Path:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel support requires openpyxl. Run: python -m pip install -r requirements.txt") from exc
    out=destination_dir/f"{safe_name(name)}.csv"
    wb=load_workbook(source,read_only=True,data_only=True)
    try:
        ws=wb.active
        with out.open("w",newline="",encoding="utf-8-sig") as fh:
            writer=csv.writer(fh)
            for row in ws.iter_rows(values_only=True):
                writer.writerow(["" if v is None else v for v in row])
    finally:
        wb.close()
    return out

def fetch_http(spec:dict,destination_dir:Path)->Path:
    r=requests.get(spec["url"],headers=_headers(spec),timeout=int(spec.get("timeout_seconds",120)),stream=True)
    r.raise_for_status()
    suffix=Path(urlparse(spec["url"]).path).suffix.lower()
    ctype=r.headers.get("content-type","").lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        suffix=".json" if "json" in ctype else ".csv"
    out=destination_dir/f"{safe_name(spec['name'])}{suffix}"
    with out.open("wb") as fh:
        for chunk in r.iter_content(1024*1024):
            if chunk:fh.write(chunk)
    return out

def fetch_s3(spec:dict,destination_dir:Path)->Path:
    try:import boto3
    except ImportError as exc:raise RuntimeError("Install boto3 for S3.") from exc
    p=urlparse(spec["uri"])
    if p.scheme!="s3" or not p.netloc or not p.path.strip("/"):
        raise ValueError("Use s3://bucket/key.ext")
    out=destination_dir/f"{safe_name(spec['name'])}{Path(p.path).suffix or '.parquet'}"
    boto3.client("s3").download_file(p.netloc,p.path.lstrip("/"),str(out))
    return out

def fetch_azure(spec:dict,destination_dir:Path)->Path:
    uri=spec["uri"]
    if uri.startswith("https://"):
        copy=dict(spec);copy["url"]=uri
        return fetch_http(copy,destination_dir)
    try:from azure.storage.blob import BlobServiceClient
    except ImportError as exc:raise RuntimeError("Install azure-storage-blob for Azure.") from exc
    conn=os.getenv("AZURE_STORAGE_CONNECTION_STRING")
    if not conn:raise RuntimeError("AZURE_STORAGE_CONNECTION_STRING is not configured.")
    p=urlparse(uri)
    client=BlobServiceClient.from_connection_string(conn).get_blob_client(p.netloc,p.path.lstrip("/"))
    out=destination_dir/f"{safe_name(spec['name'])}{Path(p.path).suffix or '.parquet'}"
    with out.open("wb") as fh:fh.write(client.download_blob().readall())
    return out

def fetch_database(spec:dict,destination_dir:Path)->Path:
    try:
        import pandas as pd
        from sqlalchemy import create_engine
    except ImportError as exc:raise RuntimeError("Install pandas and SQLAlchemy for databases.") from exc
    query=spec["query"].strip()
    if not query.lower().startswith(("select","with")):
        raise ValueError("Only read-only SELECT/CTE queries are allowed.")
    env=spec.get("connection_env")
    url=os.getenv(env) if env else spec.get("connection_url")
    if not url:raise RuntimeError(f"Database connection secret {env or 'connection_url'} is not configured.")
    out=destination_dir/f"{safe_name(spec['name'])}.csv"
    pd.read_sql_query(query,create_engine(url)).to_csv(out,index=False)
    return out

def materialize_source(spec:dict,uploaded=None,work_dir:str|Path="data/incoming")->Path:
    work_dir=Path(work_dir);work_dir.mkdir(parents=True,exist_ok=True)
    kind=spec["type"]
    if kind=="File upload":
        if uploaded is None:raise ValueError(f"No file supplied for {spec['name']}.")
        suffix=Path(uploaded.name).suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise ValueError("Supported formats: CSV, Excel (.xlsx/.xlsm), JSON, JSONL, Parquet.")
        raw=save_upload(uploaded,work_dir/f"{safe_name(spec['name'])}{suffix}")
        return excel_to_csv(raw,work_dir,spec["name"]) if suffix in {".xlsx",".xlsm"} else raw
    if kind=="API / HTTP":return fetch_http(spec,work_dir)
    if kind=="Amazon S3":return fetch_s3(spec,work_dir)
    if kind=="Azure Blob / ADLS":return fetch_azure(spec,work_dir)
    if kind=="Database":return fetch_database(spec,work_dir)
    if kind=="Local path":
        p=Path(spec["path"]).expanduser().resolve()
        if not p.exists():raise FileNotFoundError(p)
        return excel_to_csv(p,work_dir,spec["name"]) if p.suffix.lower() in {".xlsx",".xlsm"} else p
    raise ValueError(f"Unsupported source type: {kind}")
